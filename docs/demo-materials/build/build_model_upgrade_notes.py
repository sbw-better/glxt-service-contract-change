"""同步模型升级的文档共用说明和局部修订工具，保留已有段落及字体样式。"""
from pathlib import Path
from docx import Document
from docx.shared import RGBColor
import argparse

NOTES = [
    '当前历史样本约几百条且模型升级低频，普通导入和模型升级均同步执行，不使用后台执行器，也不返回异步受理202。MODEL_VERSION按完整兼容标识比较，同版本实际模型名称或维度变化时拒绝导入及识别，要求更新版本标识。',
    '升级时合并当前有效历史样本与本次Excel，按规范化Hash去重，同标签优先Excel原文、标签冲突整批拒绝。人工停用及历史失效记录不自动恢复；源文本、Hash、标签、版本异常时拒绝重建。所有合并样本都重新调用新模型，旧向量不复制。',
    '全部新向量生成成功后，一个事务插入新版本、停用旧有效行并切换有效版本。唯一约束为MODEL_VERSION与TEXT_HASH组合，旧ID、正文、标签、向量及创建时间保留。容量只计当前版本有效样本，默认10000条；Excel单次1000行限制保持。',
    '只新增HT_SAMPLE_LIBRARY一张控制表，保存有效模型身份、修订号、操作ID、状态及待恢复样本数。状态为IDLE、RUNNING、DB_COMMITTED；所有导入和人工重载通过控制行互斥，模型调用在事务外。重建成功响应code=200，含rebuild、operationId及全量入库统计。索引状态包含有效版本、配置版本、操作状态及recognitionAvailable，不提供历史任务查询。',
    '提交前失败整批回滚，旧有效版本不变并释放占用；提交后索引失败保持DB_COMMITTED并暂停识别，通过重载或重启校验数量并恢复，不再生成或重复插入。进程提交前中断时重启释放RUNNING占用；提交后中断则加载新索引后回到IDLE。精确匹配也不能绕过模型兼容检查，合同比对保留基础差异并标记识别失败。',
    '仍按单实例部署，模型配置修改后重启生效；不提供热切换或多实例广播，不保存历史重建任务。本次不执行已有库迁移或清库。请求超时后先查索引状态及有效版本，确认结果再决定重试。向量CLOB JSON、余弦相似度、Top K、加权投票及合同级汇总规则不变。详细字段和响应示例见technical-guides下的模型升级同步重建说明。',
]

REPLACEMENTS = {
    'multipart file及UserId请求头': 'JSON fileGetPath及userId',
    'multipart file；UserId头必填': 'JSON fileGetPath及userId必填',
    'multipart file、UserId请求头': 'JSON fileGetPath、userId',
    'multipart file、UserId头必填': 'JSON fileGetPath、userId必填',
    '导入和预测必须携带 UserId': '导入JSON传入userId，预测携带UserId请求头',
    '文件使用正常样本 Excel，Header 带 UserId': 'JSON传入服务器Excel路径fileGetPath及userId',

    '一张业务表、一条Sequence、一个普通索引': '历史样本表及一张控制表、一条Sequence和相关索引',
    '历史样本表及两张控制/审计表': '历史样本表及一张控制表',
    '另有控制和重建审计表': '另有一张控制表',
    '控制/审计SQL': '控制SQL',
    '导入、预测、索引及重建记录查询接口定义。': '导入、预测及索引接口定义。',
    '保留必要的版本控制和同步重建审计表。': '保留一张版本控制表。',
    '保留版本控制和同步重建审计表。': '保留一张版本控制表。',
    '第一版只新增TPIF_HTDLYB一张业务表，不依赖TPIF_XBGHTZJ。': '历史样本使用TPIF_HTDLYB，另有一张控制表，不依赖TPIF_XBGHTZJ。',
    '导入方法在单实例内使用synchronized，同一时刻只处理一份Excel。': '导入和重载通过数据库控制行占用互斥，同一时刻仅一个写入操作。',
    '在维护窗口使用原历史Excel重新导入全部样本。': '修改配置并重启后导入Excel，合并当前有效历史并同步全量重建。',
    '当前默认batch-size=1，以避免超过平台尚未确认的批量上限。': '当前默认batch-size=16，批量上限为16；输入为字符串数组。',
    'READY、EMPTY，以及仍有可用样本的DEGRADED：UP。': 'READY、EMPTY或可用DEGRADED且recognitionAvailable=true时UP；不兼容或待恢复时DOWN。',
    '如果数据库已经提交但索引重载失败，success仍为true，indexReloaded为false。': '普通增量导入提交后索引失败返回success=true、indexReloaded=false；升级重建失败返回503及DB_COMMITTED。',
    '所有业务SQL位于ContractParagraphMapper.xml，Mapper接口不使用@Select等SQL注解。': '样本及控制SQL位于对应Mapper XML，接口不使用@Select等SQL注解。',
    '四个业务接口定义。': '导入、预测及索引接口定义。',
    '全部业务SQL。': '历史样本读写SQL；控制SQL见SampleLibraryMapper.xml。',
    '分批查询Oracle全部Hash': '按当前版本分批查询Oracle Hash',
    'EMBEDDING_BATCH_SIZE默认1。': 'EMBEDDING_BATCH_SIZE默认16，上限16。',
    '不引入向量数据库、Rerank、异步任务和额外业务表。': '不引入向量数据库、Rerank或后台执行器；保留一张版本控制表。',
    '如果索引刷新失败，导入仍返回success=true但indexReloaded=false': '普通增量导入索引刷新失败返回success=true及indexReloaded=false；升级重建索引失败返回503并保留DB_COMMITTED',
    '事务提交后重建内存索引；重建失败保留旧索引并返回indexReloaded=false。': '普通增量导入提交后刷新索引，失败仅保留兼容旧快照并返回indexReloaded=false；升级重建索引失败返回503，等待恢复。',
    '查询或构建发生未捕获异常时旧索引继续服务。': '查询或构建异常时仅兼容的旧快照可继续服务；模型不兼容或升级已提交时识别不可用。',
    '默认逐条调用模型网关': '默认每批最多16条调用模型网关',
    '历史物理记录总数最多10000条': '当前版本有效样本最多10000条',
    '物理样本总数默认10000，包含停用与旧模型记录。': '当前版本有效样本默认最多10000条，失效记录不占容量。',
    'synchronized只串行化当前JVM的导入，不是跨实例分布式锁。': '导入和重载通过数据库控制行互斥，仍按单实例部署。',
    '相同标签的旧模型向量更新原行': '模型版本变化时合并有效历史与Excel，全量生成新版本行',
    '唯一Hash只保留一行，模型升级重新导入会更新该行，不同时保存多版本向量。': '同版本Hash唯一；模型升级新增新版本行，旧有效行整体停用并保留追溯。',
    '使用同一Excel重新导入会更新旧模型向量，然后再验收索引数量。': '导入Excel会合并当前有效历史并同步全量重建新版本，然后验收索引数量。',
    '模型升级后设置新的MODEL_VERSION和实际维度，并使用原历史Excel全量重新导入；不得混用不同模型版本向量。': '模型升级后更新完整MODEL_VERSION及实际模型配置并重启，再导入Excel同步重建；无需重新提供所有历史Excel，不得混用向量版本。',
    'TEXT_HASH唯一约束保证同一规范化段落只有一条物理记录。': 'MODEL_VERSION与TEXT_HASH复合唯一约束保证同版本同一规范化段落只有一条记录。',
    '相同Hash、相同类型的旧记录会更新向量、维度、模型版本并重新启用。': '同版本同Hash同标签停用记录重新启用；模型升级新增新版本行，不覆盖旧记录。',
    '累计物理记录最多10000条': '当前版本有效样本最多10000条',
    '新增后物理记录将超过10000条': '新增或重新启用后有效样本超过10000条',
    '拒绝导入；更新旧记录不增加物理记录数。': '拒绝导入；历史失效记录不占有效样本容量。',
    '历史样本物理记录总上限。': '当前版本有效样本数量上限。',
    '规范化段落SHA-256；唯一约束。': '规范化段落SHA-256；与MODEL_VERSION组合唯一。',
    '规范化文本SHA 256，全局唯一索引': '规范化文本SHA 256，与MODEL_VERSION组合唯一',
    '建立全局唯一身份': '建立版本内唯一身份',
    '第一版历史样本总数上限': '当前版本有效样本数量上限',
    '更新原ID，重生成向量并启用': '同版本停用重启保留ID；升级新增版本行',
    '重新生成向量、更新原记录并启用，计入updated。': '同版本停用重启计入updated；升级全量新增新版本记录。',
}

def replace_text(paragraph, old, new):
    # 跨run定位，仅修改命中的字符范围，保留其它run和段落属性。
    text=paragraph.text
    position=0
    while old in text[position:]:
        start=text.index(old,position); end=start+len(old); offset=0; first=True
        for run in paragraph.runs:
            part=run.text; left=offset; right=offset+len(part); offset=right
            if right<=start or left>=end: continue
            prefix=part[:max(0,start-left)]; suffix=part[max(0,end-left):] if end<right else ''
            run.text=prefix+(new if first else '')+suffix; first=False
        text=paragraph.text
        position=start+len(new)

def append_model_upgrade_notes(doc):
    for i,p in enumerate(doc.paragraphs):
        if p.text=='模型升级同步重建':
            for paragraph,text in zip(doc.paragraphs[i+1:i+1+len(NOTES)],NOTES):
                if paragraph.text != text: replace_text(paragraph,paragraph.text,text)
            return
    heading=doc.add_paragraph('模型升级同步重建','Heading 1')
    heading.runs[0].font.color.rgb=RGBColor(0,0,0)
    for text in NOTES: doc.add_paragraph(text)

def update_document(path):
    doc=Document(path)
    paragraphs=list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells: paragraphs.extend(cell.paragraphs)
            if len(row.cells)>1:
                key=row.cells[0].text.strip()
                if key=='EMBEDDING_BATCH_SIZE':
                    for p in row.cells[1].paragraphs:
                        if p.text.strip()=='1': replace_text(p,'1','16')
                if key=='SEARCH_MAX_PARAGRAPH_LENGTH':
                    for p in row.cells[1].paragraphs:
                        if p.text.strip()=='480': replace_text(p,'480','2000')
    for p in paragraphs:
        for old,new in REPLACEMENTS.items():
            if old in p.text: replace_text(p,old,new)
    append_model_upgrade_notes(doc)
    if path.name.startswith('合同变更识别完整流程与算法说明20261008'):
        for p in doc.paragraphs:
            if p.text=='模型升级同步重建': p.paragraph_format.page_break_before=True
    doc.save(path)

if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('paths',nargs='+',type=Path)
    for path in parser.parse_args().paths: update_document(path)
