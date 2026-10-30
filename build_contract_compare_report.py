from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path('合同比对逻辑现存问题与改进建议.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.52); sec.bottom_margin = Inches(.52)
sec.left_margin = Inches(.65); sec.right_margin = Inches(.65)
styles = doc.styles
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=styles[name]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
styles['Normal'].font.size=Pt(10); styles['Normal'].paragraph_format.space_after=Pt(4)
styles['Title'].font.size=Pt(22); styles['Title'].font.bold=True; styles['Title'].font.color.rgb=RGBColor(31,56,82)
title_ppr=styles['Title']._element.pPr
if title_ppr is not None:
    title_bdr=title_ppr.find(qn('w:pBdr'))
    if title_bdr is not None: title_ppr.remove(title_bdr)
styles['Heading 1'].font.size=Pt(14); styles['Heading 1'].font.bold=True; styles['Heading 1'].font.color.rgb=RGBColor(31,78,121)

p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('合同比对逻辑现存问题与改进建议')
# Remove inherited Word Title bottom border; title separation uses whitespace and typography.
ppr=p._p.get_or_add_pPr(); bdr=ppr.find(qn('w:pBdr'))
if bdr is not None: ppr.remove(bdr)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('问题梳理与改进方向'); r.font.size=Pt(10); r.font.color.rgb=RGBColor(100,100,100)

p=doc.add_paragraph()
p.add_run('结论。').bold=True
p.add_run('当前结果与生产人工勾选不一致，可能同时表现为变化类型漏报、多报、误报，以及业务类型预测不准。问题不只来自模型：文档结构解析、段落与条款匹配、历史知识库质量、投票规则和配置阈值都会影响结果。应先区分“文本差异是否正确”和“业务类型是否正确”，用有人工结论的样本逐层验证。')

def bullet(head, body):
    p=doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after=Pt(3)
    p.add_run(head).bold=True; p.add_run(body)

# compact table

doc.add_heading('主要问题与处理建议', level=1)
t=doc.add_table(rows=1, cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Light Shading Accent 1'
headers=['问题表现','可能原因','建议处理']
for c,txt in zip(t.rows[0].cells,headers):
    c.text=txt; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in c.paragraphs[0].runs: run.bold=True
rows=[
('与人工勾选相比少类型或多类型','人工标注口径不统一；一个变化段落可对应多个业务类型，但系统的候选、支持数和置信等级规则可能过滤或放大标签。','明确“一段多标签”规则；复核分歧样本；分别统计漏报、误报和多标签准确率，按业务类型调阈值。'),
('变化文本误报为删除/新增/修改','Word 的真实段落、手动换行、列表编号、表格与修订标记可能改变解析结构；解析边界与条款匹配失败会造成相同文字未配对。近期案例已发现 w:cr 与独立段落的差异，以及数字开头长正文可能被当成子条款的风险。','保留原始段落索引和边界日志；把手动换行和段落换行规范到可比较单元；限制标题启发式误识别；用真实 DOCX 建回归集并检查最终 JSON。'),
('业务类型预测不准或无可靠匹配','历史知识库覆盖不足、类型标错、相似文本重复或过时；向量相似度只表示文本接近，不保证业务标签一致。','按类型补充正例、反例和边界案例；清理重复/冲突标签；定期复核知识库并记录版本，观察各类型覆盖率与混淆情况。'),
('参数变化后结果不稳定','相似度门槛、候选数、投票支持数、候选等级门槛等配置相互作用；沿用统一参数可能不适合不同类型或文本长度。','用带人工结论的留出集做网格评估；按类型/场景调参并版本化；同时报告精确率、召回率和人工复核比例，避免只追求单一指标。'),
('上线后问题难定位','解析、匹配、预测链路可能只保留最终结果，缺少可解释证据；模型、知识库和参数版本未必与结果一起记录。','记录文档哈希、解析版本、参数快照、候选样本及分数、投票过程和最终结果；建立误报/漏报反馈闭环，回归验证后再调整。')]
for row in rows:
    cells=t.add_row().cells
    for c,txt in zip(cells,row):
        c.text=txt; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        for p in c.paragraphs:
            p.paragraph_format.space_after=Pt(2)
            for run in p.runs: run.font.size=Pt(8.3)

# widths
for row in t.rows:
    for c,w in zip(row.cells,[Inches(1.35), Inches(2.55), Inches(2.65)]): c.width=w

doc.add_heading('建议的落地顺序', level=1)
for head,body in [
('先建立对照集。','从生产人工勾选结果抽取代表性合同和段落，统一类型定义，保留人工复核记录；覆盖漏报、多报、预测错误和格式复杂文件。'),
('分层定位错误。','第一层核对原文抽取及条款归属；第二层核对变化文本配对；第三层核对业务类型候选与投票。避免用分类参数掩盖文本比对错误。'),
('再调知识库与参数。','先清理标签和样本，再在留出集上调相似度、投票和支持数参数；每次变更记录版本并比较分类型指标。'),
('持续回归与监控。','将已确认问题变成固定测试样本；对误报、漏报、无可靠匹配和人工改判持续统计，确保修复没有损害其他条款结构。')]:
    p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(3); p.add_run(head).bold=True; p.add_run(body)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(6)
r=p.add_run('说明：'); r.bold=True
p.add_run('上述原因是基于当前逻辑和已发现案例列出的风险项，知识库覆盖、人工标注一致性及具体阈值效果仍需通过生产对照集验证，不应仅凭单个案例推定。')

# Explicitly apply one font to every generated run (including table cells and headings).
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        run.font.name = 'Microsoft YaHei'
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'), 'Microsoft YaHei')
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), 'Microsoft YaHei')
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:cs'), 'Microsoft YaHei')
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = 'Microsoft YaHei'
                    rpr = run._element.get_or_add_rPr()
                    rfonts = rpr.get_or_add_rFonts()
                    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
                        rfonts.set(qn('w:' + key), 'Microsoft YaHei')

doc.core_properties.title='合同比对逻辑现存问题与改进建议'
doc.core_properties.subject='合同比对差异识别、业务类型预测和改进方向'
doc.save(out)
print(out.resolve())

